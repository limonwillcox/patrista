#!/usr/bin/env python3
"""Update sweep-titles-english.xlsx and sweep-titles-english-deepsearch.xlsx for Batch 1.

Marks all extracted Augustine treatises as 'attached' with their new repository paths,
and marks modern-only translations (FOTC 16) as 'rights-unclear'.
"""
from __future__ import annotations

import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

# Work IDs for Batch 1 extracted treatises + existing shelf files
AUGUSTINE_FILES = {
    # 16 Batch 1 treatises
    "34C1F0C7BDA4414AB336B96A98C7511D": "On Faith and the Creed.txt",
    "F08235C2B4274343A007FD1E5025D3CA": "Concerning Faith of Things Not Seen.txt",
    "5B1600C68E2046C184767CB21D079B39": "On the Profit of Believing.txt",
    "B29D3E2CD708438386D15119D200F484": "On the Creed - A Sermon to the Catechumens.txt",
    "8484FA04292246978F5D423F75CD722F": "On Continence.txt",
    "723CCB5131AB449C97379DC4A07BF0C2": "On the Good of Marriage.txt",
    "C6CB5EBA7FCB4BE69EFA247473A0DCB2": "Of Holy Virginity.txt",
    "2CF7884BD4434BDFA39F67B634AE0888": "On the Good of Widowhood.txt",
    "A3CB0ADA679E4B189C238D62AF5F1DA1": "On Lying.txt",
    "DCE35E22217B430C9BABA7D74BE1F0A3": "To Consentius Against Lying.txt",
    "8E0734A7756841F6A0103F411154AF12": "Of the Work of Monks.txt",
    "530A32B2EAEB4111BAC5D6EBE5287458": "On Patience.txt",
    "4F1E3A7FBC4D4D998974644005E04B77": "On Care to be Had for the Dead.txt",
    "8A3AAEEFE07E478882EFA4EBE0643919": "On the Perfection of Mans Righteousness.txt",
    "46FD9DEF47C641189058123996DD7E7B": "On the Proceedings of Pelagius.txt",
    "168976DED43941D2897560160A46FD40": "On Marriage and Concupiscence.txt",
    # Additional Augustine shelf files
    "0D11ECFDA83849CFB1370F81BB81A4DA": "Of the Morals of the Catholic Church.txt",
    "70BEA5798B724BB997476371A827DE43": "Letters.txt",
    "4A8B74E61A6746E8A992A4B638A6744A": "The Harmony of the Gospels.txt",
    "831B117C7BED4170A518DC52BBE7F3EE": "Our Lord's Sermon on the Mount.txt",
    "330AAEC647754C33948937FC31F6F8E5": "Homilies on the Gospel of John.txt",
    "80DDD6E4E0B24C96A0E480A8B3E687BF": "Homilies on the First Epistle of John.txt",
    "D617384F359945178EF3224244373B8D": "Expositions on the Psalms.txt",
    "A41BE7441FC0490FADB6303AB107B439": "Sermons on Selected Lessons of the New Testament.txt",
    "659DA9FACECF4E70993A5E12964A3E55": "Concerning Two Souls.txt",
    "A375161E12364B19846287C259E774D3": "Against Fortunatus.txt",
    "4C650FC41B3B44669AC764930091DAB4": "Against the Epistle of Manichaeus Called Fundamental.txt",
    "ED338925AEC5413192EB2FC79F806212": "Reply to Faustus the Manichaean.txt",
    "FBAF648DA2F74E2C8D82D45103DEEC9D": "Concerning the Nature of Good.txt",
    "B2D34707FEF94B9CB92195106A387444": "On Baptism Against the Donatists.txt",
    "F9E0134BFE3F4510AA7C9C1F801EC803": "Answer to Letters of Petilian.txt",
    "4975770998E54130B1FD3A1B10F4EB76": "On the Merits and Forgiveness of Sins.txt",
    "7014CBC08BBB46E5BABE048EFC3D1B6F": "On the Spirit and the Letter.txt",
    "85549D7A4EC94E339E1A45D2EAB1EDBC": "On Nature and Grace.txt",
    "28001CA684624675B59B3885D99FC5DF": "On the Soul and its Origin.txt",
    "4FD63CF8087346A19D13BF61DE93FAE7": "Against Two Letters of the Pelagians.txt",
    "BD484CA77DAC4A84896C0CFE1449C3DD": "On the Grace of Christ and Original Sin.txt",
    "E927859B672E4544BF52ED8915012598": "On Grace and Free Will.txt",
    "C0B95AB91A71484E85FC7A3C99930212": "On Rebuke and Grace.txt",
    "F9284C9E30044F6D8AB5E4B773097F85": "On the Predestination of the Saints.txt",
    "214473A9C90C4C8AA3B62BC57A6AB324": "On the Gift of Perseverance.txt",
}

MODERN_SERMONS = {
    "CAB214B16F4A42558067073F060B0D8D": (
        "rights-unclear",
        "FOTC 16 (1952) / WSA III/11 (1997)",
        "Modern copyright only. No pre-1928 public domain English translation exists.",
    ),
    "48B38C64BEA343C2B9F27D94E4A844C9": (
        "rights-unclear",
        "FOTC 16 (1952) / WSA III/11 (1997)",
        "Modern copyright only. No pre-1928 public domain English translation exists.",
    ),
}


def update_main() -> int:
    path = ROOT / "sweep-titles-english.xlsx"
    wb = openpyxl.load_workbook(path)
    sheet = wb.active

    headers = [c.value for c in sheet[1]]
    wid_col = headers.index("Work ID") + 1
    status_col = headers.index("Status") + 1
    note_col = headers.index("Note") + 1
    source_col = headers.index("Source") + 1

    updated = 0
    for r in range(2, sheet.max_row + 1):
        wid = str(sheet.cell(row=r, column=wid_col).value or "")
        if wid in AUGUSTINE_FILES:
            fname = AUGUSTINE_FILES[wid]
            sheet.cell(row=r, column=status_col, value="attached")
            sheet.cell(row=r, column=source_col, value=f"Fathers/English/Augustine_English/{fname}")
            sheet.cell(
                row=r,
                column=note_col,
                value=f"Extracted 1:1 public domain English text from NPNF to Fathers/English/Augustine_English/{fname} and registered in works.jsonl.",
            )
            updated += 1
        elif wid in MODERN_SERMONS:
            st, src, nt = MODERN_SERMONS[wid]
            sheet.cell(row=r, column=status_col, value=st)
            sheet.cell(row=r, column=source_col, value=src)
            sheet.cell(row=r, column=note_col, value=nt)
            updated += 1

    wb.save(path)
    return updated


def update_deep() -> int:
    path = ROOT / "sweep-titles-english-deepsearch.xlsx"
    wb = openpyxl.load_workbook(path)
    sheet = wb.active

    headers = [c.value for c in sheet[1]]
    wid_col = headers.index("Work ID") + 1
    status_col = headers.index("Status") + 1
    note_col = headers.index("Note") + 1
    source_col = headers.index("Source") + 1
    deep_st_col = headers.index("Deep Status") + 1
    deep_note_col = headers.index("Deep Analysis Note") + 1

    updated = 0
    for r in range(2, sheet.max_row + 1):
        wid = str(sheet.cell(row=r, column=wid_col).value or "")
        if wid in AUGUSTINE_FILES:
            fname = AUGUSTINE_FILES[wid]
            sheet.cell(row=r, column=status_col, value="attached")
            sheet.cell(row=r, column=source_col, value=f"Fathers/English/Augustine_English/{fname}")
            sheet.cell(
                row=r,
                column=note_col,
                value=f"Extracted 1:1 public domain English text from NPNF to Fathers/English/Augustine_English/{fname} and registered in works.jsonl.",
            )
            sheet.cell(row=r, column=deep_st_col, value="attached")
            sheet.cell(
                row=r,
                column=deep_note_col,
                value=f"Attached 1:1 English text under Fathers/English/Augustine_English/{fname}. Verified public domain NPNF translation.",
            )
            updated += 1
        elif wid in MODERN_SERMONS:
            st, src, nt = MODERN_SERMONS[wid]
            sheet.cell(row=r, column=status_col, value=st)
            sheet.cell(row=r, column=source_col, value=src)
            sheet.cell(row=r, column=note_col, value=nt)
            sheet.cell(row=r, column=deep_st_col, value=st)
            sheet.cell(row=r, column=deep_note_col, value=nt)
            updated += 1

    wb.save(path)
    return updated


def main() -> None:
    m = update_main()
    d = update_deep()
    print(f"Updated {m} rows in sweep-titles-english.xlsx")
    print(f"Updated {d} rows in sweep-titles-english-deepsearch.xlsx")


if __name__ == "__main__":
    main()

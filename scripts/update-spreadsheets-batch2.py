#!/usr/bin/env python3
"""Update sweep-titles-english.xlsx and sweep-titles-english-deepsearch.xlsx for Batch 2 (Athanasius).

Marks all extracted and attached Athanasius treatises as 'attached' with their new repository paths.
"""
from __future__ import annotations

import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

# Mapping from work_id -> extracted filename
ATHANASIUS_FILES: dict[str, str] = {
    "34C9C621F6014505B52E8D9D127C9671": "Against the Heathen.txt",
    "4D43DDE8311F44EBB1751FEDF9AC3DA0": "On the Incarnation of the Word.txt",
    "5C4F2C28A0C64D0C8BBC2317A84D7067": "To the Bishops of Egypt and Libya.txt",
    "23B12A4DF60F4DE98D7E8380625BBD3F": "Discourses Against the Arians.txt",
    "E574C1697F364586A819A6E0E882289B": "Letter to Epictetus.txt",
    "9D9A1684FC554585B3065A6F224AE048": "Letter to Adelphius.txt",
    "6EF3ED9EE1BD4312BD303315B7EA5903": "On Luke X. 22.txt",
    "4BF4E2DBCEA543D3B70FCA5C62BDF90E": "Letter to Maximus.txt",
    "0FBB951E609D4F0AAEEE0B515DE99B1A": "Life of Antony.txt",
    "195058C7795E4813AB5989A9F729F8B9": "Festal Letters and Index.txt",
    "55A83841C5CA40268745C53280BD3DF8": "First Letter to Orsisius.txt",
    "2D562C25BC6E493691BC7BFF2BE688DA": "Second Letter to Orsisius.txt",
    "12D01FC85D0C4BE69A523FAE978E0BEC": "Letters.txt",
    "1AF1547A0CD445EBB03523BD718511A8": "Letter to Amun.txt",
    "52B0019B68EA48BB9FD1E3775CD74064": "Letter to Rufinianus.txt",
    "65F0B8B73EBE4D63990BD3E4141CA9B1": "First Letter to Monks.txt",
    "3568ADFC41C04DE5ADE068ECDC1E9E1E": "Historia Acephala.txt",
    "8AC47C0F6E0E45B09649D83CDC419FC9": "Defense of the Nicene Definition.txt",
    "4969B374BA164E57A64A4DCB79801636": "On the Opinion of Dionysius.txt",
    "F447B660730E4BE2B4BFDB414FA24551": "Defense of His Flight.txt",
    "3C5F838B863641D3B49601A7ED110890": "Defense Against the Arians.txt",
    "254A2583AE974E9B9AEB7E6E15B5FA2C": "Circular Letter to Bishops.txt",
    "5BA6939A5B4C4AA8B7B448292EDD6CB4": "Letter to Serapion on the Death of Arius.txt",
    "F65D40177E154D59B4DB0A060CCF59D0": "Second Letter to Monks.txt",
    "313F0ADDF2054AD7A22038C048450CA9": "History of the Arians.txt",
    "18FA5F88F7A74C1090F589D9E4280F27": "On the Councils of Ariminum and Seleucia.txt",
    "B6538F071CEA4C239332B00E7C9B7296": "Defense Before Constantius.txt",
    "89BBF2ACCB9E4E99B1519A252D23F870": "Letter to John and Antiochus.txt",
    "216CC9ED992048D2A9CECE2C21857397": "Letter to Palladius.txt",
    "5E2DD3FF84D445AB8B704AA6A1F13742": "Letter to Dracontius.txt",
    "85B7A752597C44509E5B76A0C7687BDB": "To the Bishops of Africa.txt",
    "F7C52F971FB0498FBF96802D5BF9CED7": "Tome to the People of Antioch.txt",
    "42118E2DB2F843CBAC934A3BA8ED7D15": "Letter to the Emperor Jovian.txt",
    "1BE9BF41246F45F19809C602D41F082C": "Letter to the Emperor Jovian.txt",
    "E11BE5A1D3F9470B869D3A7A9E1629E8": "Letter to Diodorus.txt",
    "F67A4F9B05EE4AA3A78D1B59E3F202EE": "Fourth Discourse Against the Arians.txt",
    "598FA845C47A4245A636C422DD7036A7": "Letters to Lucifer.txt",
    "BC5A84ACC9E5409697B8F7638B6DF3C2": "Statement of Faith.txt",
}


def update_main() -> int:
    path = ROOT / "sweep-titles-english.xlsx"
    wb = openpyxl.load_workbook(path)
    sheet = wb.active

    headers = [c.value for c in sheet[1]]
    wid_col = headers.index("Work ID") + 1
    status_col = headers.index("Status") + 1
    note_col = headers.index("Note") + 1
    source_col = headers.index("Source") + 1

    updated = 0
    for r in range(2, sheet.max_row + 1):
        wid = str(sheet.cell(row=r, column=wid_col).value or "")
        if wid in ATHANASIUS_FILES:
            fname = ATHANASIUS_FILES[wid]
            sheet.cell(row=r, column=status_col, value="attached")
            sheet.cell(row=r, column=source_col, value=f"Fathers/English/Athanasius_English/{fname}")
            sheet.cell(
                row=r,
                column=note_col,
                value=f"Extracted 1:1 public domain English text from NPNF2-04 to Fathers/English/Athanasius_English/{fname} and registered in works.jsonl.",
            )
            updated += 1

    wb.save(path)
    return updated


def update_deep() -> int:
    path = ROOT / "sweep-titles-english-deepsearch.xlsx"
    wb = openpyxl.load_workbook(path)
    sheet = wb.active

    headers = [c.value for c in sheet[1]]
    wid_col = headers.index("Work ID") + 1
    status_col = headers.index("Status") + 1
    note_col = headers.index("Note") + 1
    source_col = headers.index("Source") + 1
    deep_st_col = headers.index("Deep Status") + 1
    deep_note_col = headers.index("Deep Analysis Note") + 1

    updated = 0
    for r in range(2, sheet.max_row + 1):
        wid = str(sheet.cell(row=r, column=wid_col).value or "")
        if wid in ATHANASIUS_FILES:
            fname = ATHANASIUS_FILES[wid]
            sheet.cell(row=r, column=status_col, value="attached")
            sheet.cell(row=r, column=source_col, value=f"Fathers/English/Athanasius_English/{fname}")
            sheet.cell(
                row=r,
                column=note_col,
                value=f"Extracted 1:1 public domain English text from NPNF2-04 to Fathers/English/Athanasius_English/{fname} and registered in works.jsonl.",
            )
            sheet.cell(row=r, column=deep_st_col, value="attached")
            sheet.cell(
                row=r,
                column=deep_note_col,
                value=f"Attached 1:1 English text under Fathers/English/Athanasius_English/{fname}. Verified public domain NPNF translation.",
            )
            updated += 1

    wb.save(path)
    return updated


def main() -> None:
    m = update_main()
    d = update_deep()
    print(f"Updated {m} rows in sweep-titles-english.xlsx")
    print(f"Updated {d} rows in sweep-titles-english-deepsearch.xlsx")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""
Update sweep-titles-english.xlsx and sweep-titles-english-deepsearch.xlsx for Batch 6 (St. Gregory of Nyssa).
"""

import openpyxl

files = [
    "sweep-titles-english.xlsx",
    "sweep-titles-english-deepsearch.xlsx",
]

UPDATES_BY_WORK_ID = {
    # Gregory of Nyssa - Letters
    "1A6B57C9C43D48D0A60CCC7F52245525": {
        "status": "attached",
        "source": "NPNF2-05, pp. 527-548 (Letters I-XVIII)",
        "note": "Extracted and attached to Fathers/English/Gregory_Nyssa_English/Letters.txt (NPNF2-05 Select Letters I-XVIII).",
    },
}

for fname in files:
    print(f"Updating {fname}...")
    wb = openpyxl.load_workbook(fname)
    ws = wb.active
    
    headers = [str(c.value or "").strip() for c in ws[1]]
    status_col = headers.index("Status") + 1
    source_col = headers.index("Source") + 1
    note_col = headers.index("Note") + 1
    work_id_col = headers.index("Work ID") + 1
    
    updated_rows = 0
    for row in range(2, ws.max_row + 1):
        wid = str(ws.cell(row=row, column=work_id_col).value or "").strip()
        if wid in UPDATES_BY_WORK_ID:
            info = UPDATES_BY_WORK_ID[wid]
            ws.cell(row=row, column=status_col).value = info["status"]
            ws.cell(row=row, column=source_col).value = info["source"]
            ws.cell(row=row, column=note_col).value = info["note"]
            updated_rows += 1
            print(f"  Row {row} [{wid}]: status -> {info['status']}")
            
    wb.save(fname)
    print(f"Saved {fname} with {updated_rows} rows updated.\n")

print("Spreadsheet updates for Batch 6 complete!")

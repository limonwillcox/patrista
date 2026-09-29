#!/usr/bin/env python3
"""
Update sweep-titles-english.xlsx and sweep-titles-english-deepsearch.xlsx for Batch 5 (Cyril of Jerusalem & Gregory of Nazianzus).
"""

import openpyxl

files = [
    "sweep-titles-english.xlsx",
    "sweep-titles-english-deepsearch.xlsx",
]

UPDATES_BY_WORK_ID = {
    # Cyril of Jerusalem - Letter to Constantius on the Cross
    "1FF843064E1F4D82B6DEAA71E3D2CF2B": {
        "status": "rights-unclear",
        "source": "CUA FOTC 64 (1970, tr. McCauley/Stephenson)",
        "note": "Only modern English translation in CUA FOTC 64 (1970); not translated in NPNF2-07 (only summarized in intro); protected by modern copyright.",
    },
    # Gregory Nazianzen - Letters
    "FD92E0441D0D4FE1BF96AC69135AA39F": {
        "status": "attached",
        "source": "NPNF2-07, pp. 437-482 (tr. Browne & Swallow)",
        "note": "Extracted and attached to Fathers/English/Gregory_Nazianzen_English/Letters.txt (NPNF2-07 Select Letters Divisions I, II, III).",
    },
    # Gregory Nazianzen - Oration II (De fuga)
    "FEF2D27310904C7397FA7076B372444C": {
        "status": "attached",
        "source": "NPNF2-07, pp. 204-227 (tr. Browne & Swallow)",
        "note": "Extracted and attached to Fathers/English/Gregory_Nazianzen_English/In Defence of His Flight to Pontus (Oration II).txt (CPG-3010.2).",
    },
    "FA1860C60A8B4DDFBFB197327CF153E7": {
        "status": "attached",
        "source": "NPNF2-07, pp. 204-227 (tr. Browne & Swallow)",
        "note": "Extracted and attached to Fathers/English/Gregory_Nazianzen_English/In Defence of His Flight to Pontus (Oration II).txt (BHL-3666t).",
    },
    # Gregory Nazianzen - Oration 10 (omitted from NPNF2-07)
    "2FE1B776ECD4497F88BD510662417989": {
        "status": "rights-unclear",
        "source": "Sources Chrétiennes / CUA FOTC 107 (2003)",
        "note": "Oration 10 is omitted from NPNF2-07 (only 21 Select Orations included); only translated in modern series (SC / FOTC 107); rights unclear / copyrighted.",
    },
    # Gregory Nazianzen - Oration 36 (omitted from NPNF2-07)
    "C04ADE31F7444181AC4BE1715BA3F393": {
        "status": "rights-unclear",
        "source": "Sources Chrétiennes / CUA FOTC 107 (2003)",
        "note": "Oration 36 is omitted from NPNF2-07 (only 21 Select Orations included); only translated in modern series (SC / FOTC 107); rights unclear / copyrighted.",
    },
    # Gregory Nazianzen - Oration 42 (Farewell 150)
    "6952B2B123924032B8A92206A1FBF705": {
        "status": "attached",
        "source": "NPNF2-07, pp. 385-395 (tr. Browne & Swallow)",
        "note": "Extracted and attached to Fathers/English/Gregory_Nazianzen_English/Farewell Oration to the 150 Bishops (Oration XLII).txt.",
    },
}

for fname in files:
    print(f"Updating {fname}...")
    wb = openpyxl.load_workbook(fname)
    ws = wb.active
    
    headers = [str(c.value or "").strip() for c in ws[1]]
    status_col = headers.index("Status") + 1
    source_col = headers.index("Source") + 1
    note_col = headers.index("Note") + 1
    work_id_col = headers.index("Work ID") + 1
    
    updated_rows = 0
    for row in range(2, ws.max_row + 1):
        wid = str(ws.cell(row=row, column=work_id_col).value or "").strip()
        if wid in UPDATES_BY_WORK_ID:
            info = UPDATES_BY_WORK_ID[wid]
            ws.cell(row=row, column=status_col).value = info["status"]
            ws.cell(row=row, column=source_col).value = info["source"]
            ws.cell(row=row, column=note_col).value = info["note"]
            updated_rows += 1
            print(f"  Row {row} [{wid}]: status -> {info['status']}")
            
    wb.save(fname)
    print(f"Saved {fname} with {updated_rows} rows updated.\n")

print("Spreadsheet updates for Batch 5 complete!")
